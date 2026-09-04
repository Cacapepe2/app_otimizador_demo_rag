"""
PMTS Presentation Engine — Vivo-style MVP
========================================

Renderiza apresentações em PowerPoint com um visual inspirado nas referências
enviadas: capa institucional, visão global do rollout e status por projeto.

Princípio:
- o motor de apresentação NÃO decide regra de negócio;
- ele apenas recebe um payload estruturado e desenha os slides;
- o payload pode ser montado por qualquer script (Risk Engine, ETL, notebook).

Dependência:
    pip install python-pptx
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Any, List, Optional
import json

from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor


# =============================================================================
# CONFIG
# =============================================================================

@dataclass(frozen=True)
class Theme:
    slide_width_in: float = 13.333
    slide_height_in: float = 7.5

    purple_dark: tuple = (32, 16, 86)
    purple_mid: tuple = (80, 45, 170)
    purple_light: tuple = (168, 137, 255)
    navy: tuple = (28, 28, 58)
    gray_text: tuple = (95, 95, 110)
    gray_line: tuple = (220, 220, 235)
    white: tuple = (255, 255, 255)
    green: tuple = (87, 163, 74)
    yellow: tuple = (219, 182, 69)
    red: tuple = (193, 74, 74)

    left_margin: float = 0.45
    right_margin: float = 0.45
    top_margin: float = 0.35
    bottom_margin: float = 0.30

    title_font_pt: int = 21
    subtitle_font_pt: int = 11
    body_font_pt: int = 10
    small_font_pt: int = 8

THEME = Theme()


# =============================================================================
# HELPERS
# =============================================================================

def _rgb(t):
    return RGBColor(*t)

def load_payload(payload_or_path):
    if isinstance(payload_or_path, dict):
        return payload_or_path
    path = Path(payload_or_path)
    return json.loads(path.read_text(encoding="utf-8"))

def _add_text(
    slide,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    font_size: int,
    color=(0, 0, 0),
    bold=False,
    align=PP_ALIGN.LEFT,
    font_name="Arial",
):
    tb = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return tb

def _rect(slide, left, top, width, height, fill, line=None, radius=False, transparency=0):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(
        shape_type, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shp.fill.solid()
    shp.fill.fore_color.rgb = _rgb(fill)
    shp.fill.transparency = transparency
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = _rgb(line)
    return shp

def _footer(slide):
    _add_text(
        slide,
        "*Documento de uso interno. Conteúdo ilustrativo para o MVP de automação de apresentações.",
        0.45, 7.05, 8.5, 0.16, 7, color=THEME.gray_text
    )

def _header_bar(slide, title: str):
    _rect(slide, 0.45, 0.45, 12.45, 0.34, THEME.purple_mid)
    _add_text(
        slide, title,
        0.45, 0.45, 12.45, 0.30,
        17, color=THEME.white, bold=True, align=PP_ALIGN.CENTER
    )

def _risk_color(name: str):
    if str(name).lower().strip() in {"baixo", "risco baixo", "low"}:
        return THEME.green
    if str(name).lower().strip() in {"medio", "médio", "risco médio", "atencao", "atenção"}:
        return THEME.yellow
    return THEME.red

def _add_donut(slide, pct: float, left: float, top: float, size: float, color):
    chart_data = ChartData()
    pct = max(0.0, min(1.0, float(pct)))
    chart_data.categories = ['A', 'B']
    chart_data.add_series('Series 1', [pct, 1.0 - pct])

    chart = slide.shapes.add_chart(
        XL_CHART_TYPE.DOUGHNUT, Inches(left), Inches(top),
        Inches(size), Inches(size), chart_data
    ).chart

    chart.has_legend = False
    chart.chart_title.has_text_frame = False
    chart.plots[0].has_data_labels = False
    chart.hole_size = 68

    series = chart.series[0]
    series.points[0].format.fill.solid()
    series.points[0].format.fill.fore_color.rgb = _rgb(color)
    series.points[1].format.fill.solid()
    series.points[1].format.fill.fore_color.rgb = _rgb((225, 225, 236))
    series.points[0].format.line.fill.background()
    series.points[1].format.line.fill.background()

    _add_text(
        slide, f"{pct:.0%}",
        left, top + size*0.36, size, 0.22,
        15, color=THEME.navy, bold=True, align=PP_ALIGN.CENTER
    )

def _progress_bar(slide, pct, left, top, width, label=None, scope_text=None, color=None):
    if color is None:
        color = THEME.purple_light
    pct = max(0.0, min(1.0, float(pct)))

    if label:
        _add_text(slide, label, left, top - 0.02, width, 0.18, 9, color=THEME.navy, bold=True)

    _rect(slide, left, top + 0.16, width, 0.14, (230, 230, 240))
    _rect(slide, left, top + 0.16, max(0.02, width * pct), 0.14, color)

    if scope_text:
        _add_text(slide, f"Escopo: {scope_text}", left + 0.05, top + 0.12, width*0.70, 0.18, 8, color=THEME.white)
    _add_text(slide, f"{pct:.0%}", left + width - 0.45, top + 0.10, 0.42, 0.18, 9, color=THEME.navy, bold=True)

def _bullet_list(slide, bullets: List[str], left, top, width, height, color=(255,255,255)):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, txt in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = str(txt)
        p.level = 0
        for run in p.runs:
            run.font.name = "Arial"
            run.font.size = Pt(10)
            run.font.color.rgb = _rgb(color)
    return tb


# =============================================================================
# SLIDES
# =============================================================================

def add_cover_slide(prs: Presentation, payload: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # background blocks
    _rect(slide, 0, 0, 13.333, 7.5, THEME.purple_dark)
    _rect(slide, 0, 0, 13.333, 7.5, THEME.purple_mid, transparency=0.52)
    _rect(slide, 0, 2.05, 13.333, 1.25, THEME.navy, transparency=0.30)

    # decorative mascot-ish abstract silhouette (approximation, not logo asset)
    _rect(slide, 8.45, 1.55, 1.20, 1.20, THEME.purple_light, radius=True, transparency=0.48)
    _rect(slide, 8.85, 2.65, 0.70, 1.85, THEME.purple_light, radius=True, transparency=0.48)
    _rect(slide, 7.95, 3.05, 1.55, 0.52, THEME.purple_light, radius=True, transparency=0.48)
    _rect(slide, 8.25, 4.05, 0.55, 1.40, THEME.purple_light, radius=True, transparency=0.48)
    _rect(slide, 9.20, 4.05, 0.55, 1.40, THEME.purple_light, radius=True, transparency=0.48)

    cover = payload.get("cover", {})
    brand = cover.get("brand", "vivo")
    title = cover.get("title", "Report Executivo")
    subtitle = cover.get("subtitle", "Rollout & Projetos | Junho 2026")

    _add_text(slide, brand, 0.95, 0.95, 2.0, 0.35, 27, color=THEME.white, bold=True)
    _add_text(slide, title, 0.95, 3.08, 5.2, 0.34, 21, color=THEME.white, bold=True)
    _add_text(slide, subtitle, 0.95, 3.43, 7.4, 0.35, 19, color=THEME.white, bold=True)


def add_global_rollout_slide(prs: Presentation, payload: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header_bar(slide, payload.get("title", "Visão Global - Rollout"))

    top = payload.get("summary_cards", [])

    # Top summary panel
    _rect(slide, 0.85, 1.05, 11.70, 1.20, THEME.navy, radius=True)
    col_x = [1.15, 4.85, 8.55]

    for idx, card in enumerate(top[:3]):
        x = col_x[idx]
        label = card.get("label", f"Grupo {idx+1}")
        pct = float(card.get("pct", 0))
        escopo = card.get("escopo", "-")
        ativos = card.get("ativos", "-")
        _add_text(slide, label, x + 0.92, 1.22, 2.20, 0.18, 10, color=THEME.white, bold=True)
        _add_donut(slide, pct, x, 1.28, 0.75, THEME.purple_light)
        _add_text(slide, str(escopo), x - 0.12, 1.87, 0.55, 0.16, 9, color=THEME.white, bold=True, align=PP_ALIGN.RIGHT)
        _add_text(slide, "ESCOPO", x - 0.04, 2.03, 0.70, 0.16, 7, color=THEME.white)
        _add_text(slide, str(ativos), x + 1.52, 1.87, 0.55, 0.16, 9, color=THEME.white, bold=True)
        _add_text(slide, "ATIVOS", x + 1.42, 2.03, 0.70, 0.16, 7, color=THEME.white)

    # Regional panel
    _rect(slide, 0.85, 2.45, 11.70, 2.40, THEME.white, line=THEME.gray_line, radius=True)
    _add_text(slide, "Farol de Status por\nRegionais", 1.15, 2.72, 1.60, 0.42, 11, color=THEME.navy, bold=True)

    legend = [("Risco Baixo", THEME.green), ("Risco Médio", THEME.yellow), ("Risco Alto", THEME.red)]
    _add_text(slide, "Legenda", 1.15, 3.55, 0.85, 0.16, 8, color=THEME.gray_text, bold=True)
    yy = 3.78
    for name, color in legend:
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(1.18), Inches(yy), Inches(0.12), Inches(0.12))
        circ.fill.solid()
        circ.fill.fore_color.rgb = _rgb(color)
        circ.line.fill.background()
        _add_text(slide, name, 1.38, yy - 0.03, 1.1, 0.15, 8, color=THEME.navy)
        yy += 0.22

    regionals = payload.get("regional_status", [])
    x_positions = [(2.80, 2.72), (6.55, 2.72), (9.15, 2.72), (6.55, 3.55), (9.15, 3.55)]
    for item, (x, y) in zip(regionals[:5], x_positions):
        risk = item.get("risk", "Risco Baixo")
        color = _risk_color(risk)
        pct = float(item.get("pct", 0))
        regional = item.get("regional", "-")
        total = item.get("total", "-")
        comment = item.get("comment", "")

        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.26), Inches(0.26))
        circ.fill.solid()
        circ.fill.fore_color.rgb = _rgb(color)
        circ.line.fill.background()
        _add_text(slide, f"{pct:.0%}", x + 0.01, y + 0.04, 0.24, 0.12, 7, color=THEME.white, bold=True, align=PP_ALIGN.CENTER)

        _add_text(slide, f"{regional} - {total}", x + 0.35, y - 0.01, 1.35, 0.16, 9, color=THEME.navy, bold=True)
        _add_text(slide, comment, x + 0.35, y + 0.15, 2.15, 0.36, 7, color=THEME.gray_text)

    # Main risks panel
    _rect(slide, 0.85, 5.10, 11.70, 1.18, THEME.navy, radius=True)
    _add_text(slide, "Principais riscos", 1.23, 5.35, 2.0, 0.18, 10, color=THEME.white, bold=True)
    risks = payload.get("main_risks", [])
    _bullet_list(slide, risks[:3], 1.15, 5.58, 10.90, 0.55, color=THEME.white)

    _footer(slide)


def add_project_status_slide(prs: Presentation, payload: Dict[str, Any]):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _header_bar(slide, payload.get("title", "Status Report por Projeto"))

    sections = payload.get("sections", [])
    current_top = 1.05

    for section in sections[:3]:
        label = section.get("section", "SEÇÃO")
        scope_total = section.get("scope_total", "")
        pct = float(section.get("pct", 0))
        items = section.get("items", [])

        # section label strip
        _rect(slide, 0.70, current_top, 4.10, 0.22, THEME.navy)
        _add_text(
            slide, f"{label} (Escopo: {scope_total})",
            0.82, current_top - 0.01, 4.0, 0.16, 8, color=THEME.white, bold=True
        )

        _add_donut(slide, pct, 0.92, current_top + 0.35, 0.85, THEME.purple_light)

        # items grid
        item_positions = [
            (2.05, current_top + 0.33),
            (5.65, current_top + 0.33),
            (9.15, current_top + 0.33),
            (2.05, current_top + 1.05),
            (5.65, current_top + 1.05),
            (9.15, current_top + 1.05),
        ]

        for item, (x, y) in zip(items[:6], item_positions):
            risk = item.get("risk", "medio")
            color = _risk_color(risk)
            circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y + 0.02), Inches(0.13), Inches(0.13))
            circ.fill.solid()
            circ.fill.fore_color.rgb = _rgb(color)
            circ.line.fill.background()
            _progress_bar(
                slide,
                item.get("pct", 0),
                x + 0.18, y - 0.02, 2.70,
                label=item.get("name", "-"),
                scope_text=item.get("scope", "-"),
                color=THEME.purple_light
            )

        current_top += 2.05

    # legend
    _add_text(slide, "Legenda:", 10.8, 1.25, 0.7, 0.16, 8, color=THEME.gray_text, bold=True)
    for idx, (name, color) in enumerate([("Risco Baixo", THEME.green), ("Risco Médio", THEME.yellow), ("Risco Alto", THEME.red)]):
        cy = 1.48 + idx*0.20
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(10.83), Inches(cy), Inches(0.10), Inches(0.10))
        circ.fill.solid()
        circ.fill.fore_color.rgb = _rgb(color)
        circ.line.fill.background()
        _add_text(slide, name, 11.00, cy - 0.03, 1.0, 0.14, 7, color=THEME.gray_text)

    _footer(slide)


# =============================================================================
# ENGINE
# =============================================================================

class VivoStylePresentationEngine:
    def __init__(self, theme: Theme = THEME):
        self.theme = theme

    def build(self, payload_or_path, output_path) -> Path:
        payload = load_payload(payload_or_path)

        prs = Presentation()
        prs.slide_width = Inches(self.theme.slide_width_in)
        prs.slide_height = Inches(self.theme.slide_height_in)

        add_cover_slide(prs, payload.get("cover_slide", payload))
        if "global_rollout" in payload:
            add_global_rollout_slide(prs, payload["global_rollout"])
        if "project_status" in payload:
            add_project_status_slide(prs, payload["project_status"])

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        prs.save(output_path)
        return output_path


def build_vivo_style_presentation(payload_or_path, output_path) -> Path:
    return VivoStylePresentationEngine().build(payload_or_path, output_path)


if __name__ == "__main__":
    demo_path = Path("examples/demo_payload.json")
    out = Path("outputs/presentations/PMTS_Rollout_Report_Demo.pptx")
    build_vivo_style_presentation(demo_path, out)
    print(out)
